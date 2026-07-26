"""Render DocumentIR -> a self-contained DOCX (images embedded, transcripts
filled into audio placeholders).
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from ..models import (
    AudioIR,
    BlockIR,
    BlockType,
    DocumentIR,
    ImageIR,
    Paragraph as ParaIR,
    RichText,
)

# formats Word can embed directly
_WORD_OK = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".emf", ".wmf"}
_CONTENT_WIDTH_IN = 6.5


class DocxRenderer:
    def __init__(self, doc_ir: DocumentIR, out_path: Path, mode: str = "exact"):
        self.ir = doc_ir
        self.out_path = out_path
        self.mode = mode
        self.doc = Document()
        self.conversions: list[str] = []

    def render(self) -> Path:
        self._setup_styles()
        self._set_core_props()
        src = self.ir.source
        self.doc.add_heading(src.title, level=0)
        if src.lesson_count:
            p = self.doc.add_paragraph(src.lesson_count)
            p.runs[0].italic = True
        for block in self.ir.blocks:
            self._render_block(block)
        self.out_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(self.out_path))
        return self.out_path

    # ------------------------------------------------------------------ #
    def _setup_styles(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        for section in self.doc.sections:
            section.left_margin = Inches(0.85)
            section.right_margin = Inches(0.85)
            section.top_margin = Inches(0.85)
            section.bottom_margin = Inches(0.85)

    def _set_core_props(self) -> None:
        cp = self.doc.core_properties
        cp.title = self.ir.source.title
        if self.ir.source.course:
            cp.subject = self.ir.source.course
        cp.author = "webpage-content-extractor"

    # ------------------------------------------------------------------ #
    def _add_rich(self, rt: RichText | None, *, style: str | None = None) -> None:
        if rt is None:
            return
        for para in rt.paragraphs:
            self._add_paragraph(para, style=style)

    def _add_paragraph(self, para: ParaIR, *, style: str | None = None) -> None:
        if para.list_type == "bullet":
            p = self.doc.add_paragraph(style="List Bullet")
        elif para.list_type == "number":
            p = self.doc.add_paragraph(style="List Number")
        elif style:
            p = self.doc.add_paragraph(style=style)
        else:
            p = self.doc.add_paragraph()
        for run in para.runs:
            if run.text == "\n":
                p.add_run().add_break()
                continue
            r = p.add_run(run.text)
            r.bold = run.bold
            r.italic = run.italic
            if run.link:
                r.font.color.rgb = RGBColor(0x0B, 0x57, 0xD0)
                r.underline = True

    # ------------------------------------------------------------------ #
    def _render_block(self, block: BlockIR) -> None:
        t = block.type
        if t == BlockType.HEADING and block.heading:
            self.doc.add_heading(block.heading.text.plain, level=min(4, block.heading.level))
        elif t == BlockType.QUOTE and block.quote:
            self._add_rich(block.quote, style="Intense Quote")
        elif t == BlockType.RICH_TEXT and block.rich_text:
            self._add_rich(block.rich_text)
        elif t == BlockType.IMAGE and block.image:
            self._render_image(block.image)
        elif t == BlockType.AUDIO and block.audio:
            self._render_audio(block.audio)
        elif t == BlockType.FLASHCARDS and block.flashcards:
            self.doc.add_heading("Flashcards", level=3)
            for i, card in enumerate(block.flashcards.cards, 1):
                fp = self.doc.add_paragraph()
                fr = fp.add_run(card.front.plain.strip() or f"Card {i}")
                fr.bold = True
                self._add_rich(card.back)
        elif t == BlockType.TABS and block.tabs:
            for tab in block.tabs.tabs:
                self.doc.add_heading(tab.title, level=3)
                self._add_rich(tab.rich_text)
                for im in tab.images:
                    self._render_image(im)
        elif t == BlockType.DIVIDER:
            if self.mode == "exact":
                self.doc.add_paragraph("* * *").alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t == BlockType.BUTTON and block.button_text:
            p = self.doc.add_paragraph()
            p.add_run(f"[{block.button_text}]").bold = True
        elif block.rich_text:
            self._add_rich(block.rich_text)

    def _render_image(self, image: ImageIR) -> None:
        asset = image.asset
        if not (asset.exists and asset.source_path):
            self._placeholder(f"Missing image (source: {asset.src_ref})")
            return
        width_in = self._width_in(asset)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        if not self._embed(run, asset, width_in):
            # remove the empty paragraph and emit a placeholder instead
            p._element.getparent().remove(p._element)
            self._placeholder(f"Could not embed image: {asset.src_ref}")
            return
        if image.caption and not image.caption.is_empty():
            cap = self.doc.add_paragraph(image.caption.plain)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True

    def _width_in(self, asset) -> float:
        width_in = _CONTENT_WIDTH_IN
        if asset.width and asset.height:
            native_in = asset.width / 96.0
            if native_in > 0:
                width_in = min(_CONTENT_WIDTH_IN, native_in)
        return width_in

    def _embed(self, run, asset, width_in: float) -> bool:
        """Try a direct embed; fall back to a Pillow re-encode (python-docx's
        header parser rejects some otherwise-valid JPEGs)."""
        path = asset.source_path
        ext = path.suffix.lower()
        if ext in _WORD_OK:
            try:
                run.add_picture(str(path), width=Inches(width_in))
                return True
            except Exception:  # noqa: BLE001 - fall through to re-encode
                pass
        buf = self._reencode_png(path)
        if buf is None:
            return False
        try:
            run.add_picture(buf, width=Inches(width_in))
            self.conversions.append(f"{path.name} -> PNG (re-encoded for DOCX embed)")
            return True
        except Exception:  # noqa: BLE001
            return False

    @staticmethod
    def _reencode_png(path: Path) -> io.BytesIO | None:
        try:
            with Image.open(path) as im:
                im.seek(0)  # first frame for animated formats
                mode = "RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB"
                buf = io.BytesIO()
                im.convert(mode).save(buf, "PNG")
                buf.seek(0)
                return buf
        except Exception:  # noqa: BLE001
            return None

    def _render_audio(self, audio: AudioIR) -> None:
        dur = audio.duration_label or (
            f"{audio.duration_seconds:.0f}s" if audio.duration_seconds else "unknown"
        )
        if audio.transcript:
            head = self.doc.add_paragraph()
            hr = head.add_run(f"Audio transcript ({dur})")
            hr.bold = True
            hr.font.color.rgb = RGBColor(0x2E, 0x5B, 0x2E)
            for line in audio.transcript.split("\n"):
                line = line.strip()
                if line:
                    self.doc.add_paragraph(line)
            if audio.asset and audio.asset.output_name:
                note = self.doc.add_paragraph()
                nr = note.add_run(f"Source audio: {audio.asset.output_name}")
                nr.italic = True
                nr.font.size = Pt(9)
        else:
            reason = "no local audio matched this placeholder"
            self._placeholder(f"Audio placeholder ({dur}) — {reason}")

    def _placeholder(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.cell(0, 0)
        p = cell.paragraphs[0]
        run = p.add_run(f"⧉ {text}")
        run.italic = True
        run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)
